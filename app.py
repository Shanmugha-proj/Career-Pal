"""Career-Pal — what should I learn next? Career-Pal knows.

Run: python -m streamlit run app.py
"""
import hashlib
import json
import os
import sys

import pandas as pd
import streamlit as st

import graph
import llm
from skills import LEVEL_LABELS

if sys.platform == "win32":
    sys.stdout.reconfigure(encoding="utf-8")

CID = os.environ.get("CAREERPAL_CID", "demo")

st.set_page_config(page_title="Career-Pal · what should I learn next?", page_icon="🧭", layout="wide",
                   initial_sidebar_state="expanded")

# ----------------------------------------------------------------------------- design system

st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&display=swap');
html, body, [data-testid="stAppViewContainer"], [data-testid="stSidebar"], .stMarkdown, .stButton, .stSlider, .stSelectbox,
[data-testid="stChatInput"] * { font-family: 'Inter', -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif; }
[data-testid="stToolbar"], .stAppDeployButton, #MainMenu, footer { display: none !important; }
header[data-testid="stHeader"] { background: transparent; }
.block-container { padding-top: 1.1rem; padding-bottom: 7rem; max-width: 1240px; }
[data-testid="stSidebar"] { background: #ffffff; border-right: 1px solid #e6e8f0; }
[data-testid="stSidebar"] .block-container { padding-top: 1.5rem; }
[data-testid="stVerticalBlockBorderWrapper"] { border: 1px solid #e6e8f0 !important; border-radius: 18px !important; background: #fff;
  box-shadow: 0 1px 2px rgba(16,24,40,.04), 0 12px 32px -20px rgba(16,24,40,.22); padding: 8px 10px 4px; margin-bottom: 14px; }
[data-testid="stExpander"] { border: 1px solid #e6e8f0; border-radius: 14px; background: #fafbfe; }
[data-testid="stFileUploader"] section { border-radius: 14px; border: 1.5px dashed #c7d2fe; background: #f8f9ff; }
.stButton > button { border-radius: 10px; font-weight: 700; padding: .55rem 1rem; }
[data-testid="stChatInput"] { border-radius: 14px; border: 1px solid #c7d2fe; background: #fff; }
[data-testid="stChatMessage"] { background: #fff; border: 1px solid #e6e8f0; border-radius: 14px; padding: 10px 14px; }

.cp-hero { background:
  linear-gradient(rgba(255,255,255,.035) 1px, transparent 1px), linear-gradient(90deg, rgba(255,255,255,.035) 1px, transparent 1px),
  radial-gradient(900px 380px at 8% -30%, rgba(99,102,241,.55), transparent 60%),
  radial-gradient(700px 300px at 95% 120%, rgba(34,211,238,.25), transparent 60%),
  linear-gradient(135deg,#0b1020 0%,#141a3a 45%,#1e1b4b 100%);
  background-size: 28px 28px, 28px 28px, auto, auto, auto; color:#fff; border-radius:22px; padding:26px 32px 26px; margin-bottom:16px;
  border:1px solid rgba(255,255,255,.08); box-shadow: 0 30px 60px -40px rgba(30,27,75,.8); }
.cp-top { display:flex; align-items:center; gap:14px; }
.cp-brand { display:flex; align-items:center; gap:12px; }
.cp-name { font-weight:800; font-size:22px; letter-spacing:-.02em; line-height:1.05; }
.cp-tag { font-size:11px; color:#9aa2c7; letter-spacing:.12em; text-transform:uppercase; font-weight:700; margin-top:3px; }
.cp-logo { width:40px;height:40px;border-radius:12px;background:linear-gradient(135deg,#6366f1 0%,#8b5cf6 50%,#22d3ee 100%); display:inline-flex;
  align-items:center;justify-content:center; box-shadow:0 10px 24px -8px rgba(99,102,241,.9), inset 0 1px 0 rgba(255,255,255,.35); flex:0 0 40px; }
.cp-eyebrow { margin-top:28px; font-size:11.5px; letter-spacing:.16em; text-transform:uppercase; color:#8b93b8; font-weight:700; }
.sb-brand { display:flex; align-items:center; gap:10px; margin: -6px 0 18px; }
.sb-brand .cp-logo { width:34px; height:34px; border-radius:10px; flex-basis:34px; }
.sb-name { font-weight:800; font-size:18px; letter-spacing:-.02em; color:#0f172a; line-height:1.05; }
.sb-tag { font-size:10.5px; color:#64748b; letter-spacing:.1em; text-transform:uppercase; font-weight:700; margin-top:2px; }
.steps { display:grid; grid-template-columns: repeat(3, 1fr); gap:12px; margin-top:22px; }
.stepc { background:rgba(255,255,255,.07); border:1px solid rgba(255,255,255,.14); border-radius:14px; padding:14px 16px; }
.stepc .n { display:inline-flex; width:26px; height:26px; border-radius:8px; background:linear-gradient(135deg,#6366f1,#22d3ee); align-items:center; justify-content:center; font-weight:800; font-size:13px; margin-bottom:8px; }
.stepc .t { font-weight:700; font-size:15px; } .stepc .d { font-size:13px; color:#c7cbe0; margin-top:3px; line-height:1.45; }
.cp-pills { margin-left:auto; display:flex; gap:8px; flex-wrap:wrap; }
.pill { display:inline-flex; align-items:center; gap:7px; padding:5px 11px; border-radius:999px; font-size:12.5px; font-weight:600;
  background:rgba(255,255,255,.09); border:1px solid rgba(255,255,255,.16); color:#e5e7eb; }
.pill .dot { width:7px;height:7px;border-radius:50%; background:#34d399; box-shadow:0 0 0 3px rgba(52,211,153,.22); }
.cp-title { font-size:44px; font-weight:800; letter-spacing:-.028em; line-height:1.06; margin:10px 0 12px; }
.cp-title em { font-style:normal; background:linear-gradient(90deg,#a5b4fc,#67e8f9); -webkit-background-clip:text; background-clip:text; color:transparent; }
.cp-sub { font-size:16px; color:#c7cbe0; max-width:860px; line-height:1.55; }
.cp-sub b { color:#fff; }
.kpis { display:grid; grid-template-columns: repeat(4, 1fr); gap:12px; margin-top:22px; }
.kpi { position:relative; background:rgba(255,255,255,.07); border:1px solid rgba(255,255,255,.14); border-radius:14px; padding:14px 16px 14px 20px; min-height:88px; }
.kpi:before { content:""; position:absolute; left:0; top:14px; bottom:14px; width:3px; border-radius:3px; background:linear-gradient(#a5b4fc,#67e8f9); }
.kpi .k { font-size:11.5px; text-transform:uppercase; letter-spacing:.08em; color:#aab0cc; font-weight:700; }
.kpi .v { font-size:26px; font-weight:800; margin-top:5px; letter-spacing:-.01em; line-height:1.1; }
.kpi .v small { font-size:13px; font-weight:600; color:#c7cbe0; margin-left:6px; }
.kpi .d { font-size:13px; color:#86efac; font-weight:800; margin-left:8px; }

.sec { display:flex; align-items:center; gap:12px; margin: 8px 0 12px; }
.sec .n { width:30px;height:30px;border-radius:9px;background:#eef2ff;color:#4f46e5;display:inline-flex;align-items:center;justify-content:center;font-weight:800;font-size:14px; }
.sec h2 { margin:0; font-size:21px; font-weight:800; letter-spacing:-.01em; color:#0f172a; }
.sec .hint { margin-left:auto; color:#64748b; font-size:13px; }
.big { font-size:50px; font-weight:800; letter-spacing:-.025em; line-height:1.05; margin:4px 0 2px; color:#0f172a; }
.big em { font-style:normal; color:#4f46e5; }
.big .delta { font-size:20px; color:#059669; font-weight:800; margin-left:10px; vertical-align:middle; background:#ecfdf5; padding:4px 10px; border-radius:10px; }
.sub { color:#64748b; font-size:15px; margin:0 0 14px; }
.chips { margin: 4px 0 6px; }
.chip { display:inline-flex; align-items:center; gap:6px; padding:6px 12px; margin:4px 6px 4px 0; border-radius:999px; font-size:13.5px; font-weight:600; border:1px solid transparent; }
.l3 { background:#ecfdf5; color:#047857; border-color:#a7f3d0; } .l2 { background:#eef2ff; color:#4338ca; border-color:#c7d2fe; } .l1 { background:#f1f5f9; color:#475569; border-color:#e2e8f0; }
.chip.new { box-shadow:0 0 0 3px rgba(245,158,11,.35); background:#fffbeb; color:#b45309; border-color:#fcd34d; }
.chip .lv { font-size:11px; opacity:.8; font-weight:700; }
.jobs { display:grid; grid-template-columns: repeat(3, 1fr); gap:12px; }
.job { border:1px solid #e6e8f0; border-radius:14px; padding:14px 14px 12px; background:#fff; display:flex; flex-direction:column; gap:10px; }
.job-top { display:flex; align-items:flex-start; gap:10px; }
.avatar { width:38px;height:38px;border-radius:10px; display:inline-flex;align-items:center;justify-content:center;font-weight:800;color:#fff;font-size:14px; flex:0 0 38px; }
.job-title { font-weight:700; font-size:14.5px; line-height:1.25; color:#0f172a; }
.job-co { color:#64748b; font-size:12.5px; margin-top:2px; }
.fit { margin-left:auto; font-weight:800; font-size:18px; letter-spacing:-.01em; white-space:nowrap; }
.bar { height:6px; background:#eef0f6; border-radius:999px; overflow:hidden; }
.bar span { display:block; height:100%; border-radius:999px; }
.tags { display:flex; flex-wrap:wrap; gap:5px; }
.tag { font-size:11.5px; padding:3px 8px; border-radius:7px; font-weight:600; }
.tag.ok { background:#ecfdf5; color:#047857; } .tag.miss { background:#fff1f2; color:#be123c; } .tag.part { background:#fffbeb; color:#b45309; }
.job-foot { display:flex; align-items:center; margin-top:auto; gap:8px; }
.badge { font-size:11.5px; font-weight:700; padding:3px 8px; border-radius:7px; background:#ecfdf5; color:#047857; }
.badge.syn { background:#f1f5f9; color:#475569; }
.apply { margin-left:auto; font-size:12.5px; font-weight:700; color:#4f46e5; text-decoration:none; padding:6px 10px; border-radius:8px; background:#eef2ff; }
.apply:hover { background:#e0e7ff; }
.lb { display:flex; flex-direction:column; gap:8px; }
.lb-row { display:grid; grid-template-columns: 34px 1fr 120px 96px; gap:14px; align-items:center; padding:12px 14px; border:1px solid #e6e8f0; border-radius:12px; background:#fff; }
.lb-row.top { border-color:#c7d2fe; background:linear-gradient(90deg,#eef2ff, #fff 60%); }
.lb-rank { width:30px;height:30px;border-radius:9px;background:#f1f5f9;color:#334155;display:inline-flex;align-items:center;justify-content:center;font-weight:800; }
.lb-row.top .lb-rank { background:#4f46e5; color:#fff; }
.lb-skill { font-weight:700; font-size:15px; color:#0f172a; }
.lb-course { color:#64748b; font-size:12px; font-weight:500; margin-left:8px; }
.cl { color:#4f46e5; text-decoration:none; font-weight:600; } .cl:hover { text-decoration:underline; }
.lb-links { font-size:12.5px; margin-top:6px; color:#64748b; line-height:1.6; }
.free { font-size:10.5px; font-weight:700; color:#047857; background:#ecfdf5; border-radius:6px; padding:1px 6px; margin-left:4px; }
.lb-bar { height:8px; background:#eef0f6; border-radius:999px; overflow:hidden; margin-top:7px; }
.lb-bar span { display:block; height:100%; background:linear-gradient(90deg,#6366f1,#22d3ee); border-radius:999px; }
.lb-num { font-weight:800; font-size:20px; color:#0f172a; text-align:right; white-space:nowrap; }
.lb-num small, .lb-days small { font-size:12px; color:#64748b; font-weight:600; margin-left:4px; }
.lb-days { font-weight:700; font-size:16px; color:#334155; text-align:right; white-space:nowrap; }
.callout { display:flex; align-items:center; gap:12px; padding:14px 16px; border-radius:14px; background:linear-gradient(90deg,#eef2ff,#ecfeff); border:1px solid #c7d2fe; margin-bottom:12px; font-size:16px; color:#0f172a; }
.callout .vs { color:#64748b; font-size:14px; margin-left:auto; text-align:right; }
.path { display:flex; align-items:stretch; margin:8px 0 6px; overflow-x:auto; padding-bottom:4px; }
.step { flex:1 1 170px; min-width:170px; position:relative; padding:14px 14px 12px; border:1px solid #e6e8f0; border-radius:14px; background:#fff; }
.step + .step { margin-left:28px; }
.step + .step:before { content:""; position:absolute; left:-28px; top:50%; width:28px; height:2px; background:#c7d2fe; }
.step .day { font-size:11.5px; text-transform:uppercase; letter-spacing:.08em; color:#4f46e5; font-weight:800; }
.step .sk { font-size:17px; font-weight:800; line-height:1.2; margin:5px 0 6px; color:#0f172a; overflow-wrap:anywhere; }
.step .m { font-size:13px; color:#334155; font-weight:600; } .step .c { font-size:11.5px; color:#64748b; margin-top:4px; }
.status { font-size:12.5px; color:#334155; }
.status .row { display:flex; justify-content:space-between; gap:8px; padding:7px 0; border-bottom:1px dashed #e6e8f0; }
.status .row:last-child { border-bottom:none; }
.status .ok { color:#059669; font-weight:700; font-variant-numeric: tabular-nums; }
.status .lbl { color:#64748b; cursor: help; text-decoration: underline dotted rgba(100,116,139,.6); text-underline-offset: 3px; }
.status .row { cursor: help; }
.banner { padding:12px 16px; border-radius:12px; background:#ecfdf5; border:1px solid #a7f3d0; color:#065f46; font-weight:600; margin: 4px 0 10px; }
.note { color:#64748b; font-size:12.5px; line-height:1.5; }
@media (max-width: 1100px) { .jobs { grid-template-columns: repeat(2, 1fr);} .kpis { grid-template-columns: repeat(2, 1fr);} }
</style>
""", unsafe_allow_html=True)

LOGO_SVG = ('<svg width="20" height="20" viewBox="0 0 24 24" fill="none"><path d="M7 17.5 10.5 7.5M17 17.5 13.5 7.5M8 19h8" '
            'stroke="#fff" stroke-width="2.2" stroke-linecap="round"/><circle cx="5" cy="19" r="3" fill="#fff"/>'
            '<circle cx="19" cy="19" r="3" fill="#fff"/><circle cx="12" cy="5" r="3" fill="#fff"/></svg>')

SHORT = {
    "CI/CD (GitLab CI/GitHub Actions)": "CI/CD pipelines", "Prometheus/Grafana": "Prometheus / Grafana",
    "SRE Practices (SLOs/On-call)": "SRE practices", "Networking (L2/L3)": "Networking", "Incident Management/RCA": "Incident mgmt",
    "Cloud Security/IAM": "Cloud security", "GitOps (ArgoCD/Flux)": "GitOps", "Logging (ELK/Loki)": "Logging",
    "Bare-metal Provisioning": "Bare-metal", "Virtualization (VMware/KVM)": "Virtualization", "Storage (Ceph/SAN)": "Storage",
    "Secrets Management (Vault)": "Vault", "Service Mesh (Istio)": "Service mesh", "Telecom Core Networks": "Telecom core",
    "Technical Documentation": "Tech writing",
}
short = lambda s: SHORT.get(s, s)
AVATAR_COLORS = ["#4f46e5", "#0891b2", "#059669", "#d97706", "#db2777", "#7c3aed", "#2563eb", "#dc2626", "#0d9488"]


# ----------------------------------------------------------------------------- helpers

def read_resume_file(up) -> str:
    """Extract plain text from an uploaded PDF / DOCX / TXT."""
    import io
    import re
    name = up.name.lower()
    data = up.getvalue()
    if name.endswith(".pdf"):
        from pypdf import PdfReader
        text = "\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(data)).pages)
    elif name.endswith(".docx"):
        from docx import Document
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        parts += [c.text for t in doc.tables for row in t.rows for c in row.cells]
        text = "\n".join(parts)
    else:
        text = data.decode("utf-8", errors="replace")
    text = text.replace("\ufffd", "-").replace(" -", "-")
    return re.sub(r"[ \t]+\n", "\n", text).strip()


def load_resume_text() -> str:
    try:
        return open("demo/resume.txt", encoding="utf-8").read()
    except FileNotFoundError:
        return ""


def cached_skills() -> list[dict]:
    try:
        return json.load(open("demo/resume_skills.json"))["skills"]
    except FileNotFoundError:
        return []


def skills_from_graph() -> list[dict]:
    return graph.candidate_skills(CID)


def signature(skills: list[dict]) -> tuple:
    return tuple(sorted((s["name"], s["level"]) for s in skills))


def get_fits(sig: tuple) -> list[dict]:
    cache = st.session_state.setdefault("fits_cache", {})
    if sig not in cache:
        cache.clear()
        cache[sig] = graph.job_fits(CID)
    return cache[sig]


def get_plan(sig: tuple, cutoff: float, budget: int) -> dict:
    cache = st.session_state.setdefault("plan_cache", {})
    key = (sig, cutoff, budget)
    if key not in cache:
        cache[key] = graph.plan_path(CID, cutoff, budget_days=budget)
    return cache[key]


def level_label(skill: str, level: int) -> str:
    return LEVEL_LABELS.get(skill, {}).get(level, {1: "basic", 2: "pro", 3: "expert"}[level])


def chips(skills: list[dict], highlight: set[str] = frozenset()) -> str:
    return '<div class="chips">' + "".join(
        f'<span class="chip l{s["level"]} {"new" if s["name"] in highlight else ""}">{s["name"]}'
        f'<span class="lv">{"NEW · " if s["name"] in highlight else ""}{level_label(s["name"], s["level"])}</span></span>'
        for s in skills) + "</div>"


def fit_color(f: float) -> str:
    return "#059669" if f >= 0.8 else "#4f46e5" if f >= 0.6 else "#d97706"


def avatar(company: str) -> str:
    color = AVATAR_COLORS[int(hashlib.md5(company.encode()).hexdigest(), 16) % len(AVATAR_COLORS)]
    initials = "".join(w[0] for w in company.replace("&", "").split()[:2]).upper() or "?"
    return f'<div class="avatar" style="background:{color}">{initials}</div>'


def job_card(f: dict) -> str:
    fit = f["fit"]
    ok = [s for s in f["matched"][:4]]
    gaps = f["gaps"][:3]
    other = (f.get("other") or [])[:3]
    tags = "".join(f'<span class="tag ok">{short(s)}</span>' for s in ok) + "".join(
        f'<span class="tag {"part" if g["m"] > 0 else "miss"}">{"↑ " if g["m"] > 0 else "+ "}{short(g["s"])}</span>' for g in gaps) + "".join(
        f'<span class="tag miss">+ {o}</span>' for o in other)
    real = f.get("source") == "real"
    badge = '<span class="badge">● Real posting</span>' if real else '<span class="badge syn">synthetic</span>'
    link = f'<a class="apply" href="{f["url"]}" target="_blank">Apply ↗</a>' if f.get("url") else ""
    return (f'<div class="job"><div class="job-top">{avatar(f["company"])}<div><div class="job-title">{f["title"]}</div>'
            f'<div class="job-co">{f["company"]} · {f["location"]}</div></div>'
            f'<div class="fit" style="color:{fit_color(fit)}">{round(100 * fit)}%</div></div>'
            f'<div class="bar"><span style="width:{round(100 * fit)}%;background:{fit_color(fit)}"></span></div>'
            f'<div class="tags">{tags}</div><div class="job-foot">{badge}{link}</div></div>')


@st.cache_data
def load_courses() -> dict:
    """Best courses per skill: picked by DeepSeek V4 Flash on ai&, every link verified (python courses.py)."""
    try:
        return json.load(open("demo/courses.json", encoding="utf-8"))
    except FileNotFoundError:
        return {}


def course_links(skill: str, n: int = 2) -> str:
    import html as _h
    out = []
    import re as _re
    for c in load_courses().get(skill, [])[:n]:
        prov = _re.sub(r"\s*\(.*?\)", "", c["provider"]).strip()
        prov = prov if len(prov) <= 24 else " ".join(prov.split()[:2])
        t = c["title"]
        for name in {c["provider"], prov}:
            if t.lower().startswith(name.lower()):
                t = t[len(name):].lstrip(" :–-")
            t = _re.sub(r"\s*\(" + _re.escape(name) + r"\)\s*$", "", t)
        t = t if len(t) <= 52 else t[:50] + "…"
        out.append(f'<a class="cl" href="{_h.escape(c["url"])}" target="_blank" rel="noopener">{_h.escape(prov)}: {_h.escape(t)} ↗</a>'
                   + ('<span class="free">free</span>' if c.get("free") else ""))
    return " · ".join(out)


def apply_skill(skill: str, level: int, source: str):
    """Agent memory write + bookkeeping for the 'what changed' deltas."""
    cur = skills_from_graph()
    st.session_state["before"] = {"fits": get_fits(signature(cur)), "sig": signature(cur),
                                  "plan": st.session_state.get("last_plan")}
    graph.mark_skill(CID, skill, level, source=source)
    st.session_state["highlight"] = {skill}
    st.session_state.pop("plan_cache", None)


def section(n: int, title: str, hint: str = ""):
    st.markdown(f'<div class="sec"><span class="n">{n}</span><h2>{title}</h2><span class="hint">{hint}</span></div>',
                unsafe_allow_html=True)


# ----------------------------------------------------------------------------- sidebar

with st.sidebar:
    st.markdown("**Controls**")
    model_ids = list(llm.MODELS)
    model = st.selectbox("Model (ai& inference)", model_ids, index=model_ids.index(llm.MODEL) if llm.MODEL in model_ids else 0,
                         format_func=lambda m: llm.MODELS[m][0],
                         help="The open-weight model on ai& inference (hosted in Japan) that extracts your skills and runs the agent.")
    st.session_state["model"] = model
    cutoff_pct = st.slider("Job Description Match %", 30, 90, 50, 5, format="%d%%",
                           help="You count as a match for a job when your skills cover at least this much of its job description.")
    cutoff = cutoff_pct / 100
    budget = st.slider("How quickly you want to move in your career (days)", 14, 120, 45, 1,
                       help="How many days of learning you can invest. The plan picks the skills that unlock the most jobs per day.")
    st.divider()
    st.markdown("**Live status**")
    status = st.empty()
    st.divider()
    if st.button("Reset demo", width="stretch"):
        graph.set_candidate_skills(CID, [])
        for k in ("fits_cache", "plan_cache", "before", "highlight", "chat", "extract_secs", "extract_note", "last_plan",
                  "resume_text", "uploaded_name", "agent_secs"):
            st.session_state.pop(k, None)
        st.rerun()
    st.markdown('<p class="note">Corpus: real Tokyo postings pulled today from public job-board APIs (Greenhouse, Lever, Workable, '
                'Ashby, SmartRecruiters, Workday) and structured with ai&.</p>', unsafe_allow_html=True)

# ----------------------------------------------------------------------------- data

skills = skills_from_graph()
sig = signature(skills) if skills else ()
fits = get_fits(sig) if skills else []
qualified = [f for f in fits if f["fit"] >= cutoff]
n_q = len(qualified)
n_real = sum(1 for f in fits if f.get("source") == "real")
n_companies = len({f["company"] for f in fits})

lb_all = graph.leaderboard(CID, cutoff, limit=25) if skills else []
lb_all = [r for r in lb_all if r["unlocks"] >= 2]              # drop single-job noise
lb = lb_all[:8]
top = lb[0] if lb else None                                       # best next move = most jobs per day of learning
biggest = max(lb_all, key=lambda r: r["unlocks"]) if lb_all else None   # biggest single unlock, whatever it costs

plan = None
if skills:
    with st.spinner("Re-planning in Neo4j…"):
        plan = get_plan(sig, cutoff, budget)
    st.session_state["last_plan"] = plan

before = st.session_state.get("before")
delta_jobs = None
if before and before["sig"] != sig:
    delta_jobs = n_q - sum(1 for f in before["fits"] if f["fit"] >= cutoff)

# ----------------------------------------------------------------------------- hero

model_label = llm.MODELS.get(model, (model,))[0]
kpis = ""
if skills:
    kpi_jobs = f'<div class="kpi"><div class="k">Real postings analysed</div><div class="v">{len(fits)}<small>{n_companies} companies</small></div></div>'
    d = f'<span class="d">{"+" if delta_jobs > 0 else ""}{delta_jobs}</span>' if delta_jobs else ""
    kpi_q = f'<div class="kpi"><div class="k">You qualify for</div><div class="v">{n_q}{d}<small>≥ {cutoff_pct}% JD match</small></div></div>'
    kpi_next = (f'<div class="kpi"><div class="k">Best next move</div><div class="v">{short(top["skill"])}<small>+{top["unlocks"]} jobs · {top["days"]}d</small></div></div>'
                if top else '<div class="kpi"><div class="k">Best next move</div><div class="v">—</div></div>')
    kpi_path = (f'<div class="kpi"><div class="k">Fastest path</div><div class="v">{plan["total_days"]} days<small>→ +{plan["total_unlocked"]} jobs</small></div></div>'
                if plan and plan["steps"] else '<div class="kpi"><div class="k">Fastest path</div><div class="v">—</div></div>')
    kpis = f'<div class="kpis">{kpi_jobs}{kpi_q}{kpi_next}{kpi_path}</div>'

steps = ('<div class="steps">'
         '<div class="stepc"><span class="n">1</span><div class="t">Upload your resume</div><div class="d">PDF, DOCX or text. Processed on ai& infrastructure in Japan — it never leaves the country.</div></div>'
         '<div class="stepc"><span class="n">2</span><div class="t">Your skills become a graph</div><div class="d">Structured into a 44-skill taxonomy and written as edges into a Neo4j memory graph.</div></div>'
         '<div class="stepc"><span class="n">3</span><div class="t">See the cheapest skill to learn next</div><div class="d">Scored against 370 real Tokyo postings — and re-planned the moment you learn something.</div></div>'
         '</div>') if not skills else ""
eyebrow = f"Tokyo · {len(fits)} real postings · {n_companies} companies · pulled today" if skills else "Tokyo · 370 real postings · 199 companies · pulled today"
st.markdown(
    f'<div class="cp-hero"><div class="cp-top"><div class="cp-brand"><span class="cp-logo">{LOGO_SVG}</span>'
    f'<div><div class="cp-name">Career-Pal</div><div class="cp-tag">Skill-graph career planner</div></div></div>'
    f'<span class="cp-pills"><span class="pill"><span class="dot"></span>Neo4j AuraDB · memory graph</span>'
    f'<span class="pill"><span class="dot"></span>ai& inference · {model_label}</span>'
    f'<span class="pill">🇯🇵 Data stays in Japan</span></span></div>'
    f'<div class="cp-eyebrow">{eyebrow}</div>'
    f'<div class="cp-title">What should you learn next?<br><em>Career-Pal knows.</em></div>'
    f'<div class="cp-sub">Japan is short <b>~790,000 IT workers by 2030</b> (METI) — a <b>skills</b> gap, not a headcount gap. '
    f'Every job tool asks “which jobs fit me?”. Career-Pal inverts it: <b>what is the cheapest change to <i>you</i> that fits the most jobs?</b></div>'
    f'{kpis}{steps}</div>', unsafe_allow_html=True)

# ----------------------------------------------------------------------------- 1 · resume

with st.container(border=True):
    section(1, "Your resume", f"Structured extraction into a 44-skill taxonomy · {model_label} on ai&, hosted in Japan")
    with st.expander("Upload or paste your resume", expanded=not skills):
        up = st.file_uploader("Upload resume (PDF, DOCX or TXT)", type=["pdf", "docx", "txt"], label_visibility="collapsed")
        if up is not None and st.session_state.get("uploaded_name") != up.name:
            try:
                st.session_state["resume_text"] = read_resume_file(up)
                st.session_state["uploaded_name"] = up.name
                st.toast(f"Loaded {up.name} · {len(st.session_state['resume_text'].split())} words")
            except Exception as e:  # noqa: BLE001
                st.error(f"Could not read {up.name}: {e}")
        resume_text = st.text_area("Resume text", st.session_state.get("resume_text") or "", height=200,
                                   label_visibility="collapsed", placeholder="…or paste your resume text here")
        c1, c2 = st.columns([1, 3])
        if c1.button("Analyze with ai& ⚡", type="primary", width="stretch", disabled=not resume_text.strip()):
            note = ""
            try:
                with st.spinner(f"{model_label} · extracting skills on ai& in Japan…"):
                    extracted, secs = llm.extract_skills(resume_text, model=model)
                if not extracted:
                    raise ValueError("no skills returned")
            except Exception as e:  # noqa: BLE001 - never let a live LLM call sink the demo
                if resume_text.strip() == load_resume_text().strip() and cached_skills():
                    extracted, secs, note = cached_skills(), 0.0, f"ai& unavailable ({type(e).__name__}) — loaded cached extraction"
                else:
                    st.error(f"ai& inference failed ({type(e).__name__}: {str(e)[:120]}). Try again or switch model in the sidebar.")
                    st.stop()
            graph.set_candidate_skills(CID, extracted)
            st.session_state["extract_secs"] = secs
            st.session_state["extract_note"] = note
            for k in ("fits_cache", "plan_cache", "before", "highlight", "last_plan"):
                st.session_state.pop(k, None)
            st.rerun()
        c2.markdown('<p class="note" style="margin-top:10px">Your resume is some of the most sensitive data you own. '
                    'Every token is processed on ai& infrastructure in Japan and becomes <b>HAS</b> edges in your Neo4j memory graph — '
                    'nothing is stored as free text.</p>', unsafe_allow_html=True)

    if skills:
        st.markdown(chips(skills, st.session_state.get("highlight", set())), unsafe_allow_html=True)
        secs = st.session_state.get("extract_secs")
        meta = f"**{len(skills)} skills** in your memory graph"
        if secs:
            meta += f" · extracted in **{secs:.1f}s** by {model_label}"
        if st.session_state.get("extract_note"):
            meta += f" · ⚠️ {st.session_state['extract_note']}"
        st.caption(meta)
        st.markdown('<p class="note" style="margin-top:-6px">Levels are judged by the model from evidence in your resume: '
                    '<b>expert</b> = certified, led it, or ran it at scale · <b>pro</b> = independent professional use · '
                    '<b>basic</b> = coursework or exposure. A posting that needs a higher level than you have gives half credit (↑ tags).</p>',
                    unsafe_allow_html=True)

if not skills:
    st.markdown('<p class="note" style="text-align:center;font-size:14px;margin-top:6px">Upload a resume (PDF / DOCX / TXT) and click '
                '<b>Analyze with ai&</b>. Your skills become edges in a Neo4j memory graph, scored against 370 real Tokyo postings.</p>',
                unsafe_allow_html=True)
    st.stop()

# ----------------------------------------------------------------------------- 2 · where you stand

with st.container(border=True):
    section(2, "Where you stand", f"{n_real} real postings · {n_companies} companies · scored in Neo4j")
    d_html = f'<span class="delta">{"+" if delta_jobs > 0 else ""}{delta_jobs}</span>' if delta_jobs else ""
    st.markdown(f'<p class="big">You qualify for <em>{n_q}</em> of {len(fits)} jobs{d_html}</p>'
                f'<p class="sub">at ≥ {cutoff_pct}% job description match · Tokyo infrastructure / cloud / platform roles · '
                f'fit = weighted coverage of each posting’s required skills, explainable skill by skill</p>', unsafe_allow_html=True)
    show = sorted(qualified, key=lambda f: -f["fit"])[:9]
    if show:
        st.markdown('<div class="jobs">' + "".join(job_card(f) for f in show) + "</div>", unsafe_allow_html=True)
    else:
        st.warning("No jobs at this match % — lower the Job Description Match % in the sidebar.")
    with st.expander(f"All {n_q} qualifying postings"):
        df = pd.DataFrame([{
            "Fit": round(100 * f["fit"]), "Title": f["title"], "Company": f["company"], "Location": f["location"],
            "Link": f.get("url") or None,
            "Missing": ", ".join([f'{g["s"]}{"" if g["m"] == 0 else " ↑"}' for g in f["gaps"]] + list(f.get("other") or [])) or "—",
        } for f in sorted(qualified, key=lambda f: -f["fit"])])
        st.dataframe(df, hide_index=True, width="stretch", height=360, column_config={
            "Fit": st.column_config.ProgressColumn("Fit", format="%d%%", min_value=0, max_value=100),
            "Link": st.column_config.LinkColumn("Link", display_text="apply ↗"),
        })

# ----------------------------------------------------------------------------- 3 · learn next

with st.container(border=True):
    section(3, "Learn next — the cheapest edge, not the biggest", "ranked by jobs unlocked per day of learning")
    if top:
        vs = (f'<span class="vs">biggest single unlock: <b>{biggest["skill"]}{" " + level_label(biggest["skill"], biggest["level"]) if biggest["skill"] in LEVEL_LABELS else ""}</b>'
              f' → +{biggest["unlocks"]} jobs — but <b>{biggest["days"]} days</b> ({biggest["jobs_per_day"]:.2f} jobs/day)</span>'
              if biggest and biggest["skill"] != top["skill"] else "")
        st.markdown(f'<div class="callout">🏆&nbsp; <span>Best next move: <b>{top["skill"]}</b> → <b>+{top["unlocks"]} jobs</b> in <b>{top["days"]} days</b>'
                    f' <span style="color:#4f46e5;font-weight:800">· {top["jobs_per_day"]:.1f} jobs/day</span></span>{vs}</div>',
                    unsafe_allow_html=True)
        max_unlocks = max(r["unlocks"] for r in lb)
        rows = "".join(
            f'<div class="lb-row {"top" if i == 0 else ""}"><div class="lb-rank">{i + 1}</div>'
            f'<div><div class="lb-skill">{r["skill"]}{" (" + level_label(r["skill"], r["level"]) + ")" if r["skill"] in LEVEL_LABELS else ""}'
            f'</div>'
            f'<div class="lb-bar"><span style="width:{max(4, round(100 * r["unlocks"] / max_unlocks))}%"></span></div>'
            f'<div class="lb-links">Learn it: {course_links(r["skill"]) or (r["course"] or "")}</div></div>'
            f'<div class="lb-num">+{r["unlocks"]}<small>jobs</small></div>'
            f'<div class="lb-days">{r["days"]}<small>days</small><br><small>{r["jobs_per_day"]:.2f} / day</small></div></div>'
            for i, r in enumerate(lb))
        st.markdown(f'<div class="lb">{rows}</div>', unsafe_allow_html=True)
        st.markdown('<p class="note" style="margin-top:8px">Course picks by DeepSeek V4 Flash on ai& inference · every link checked before it is shown.</p>',
                    unsafe_allow_html=True)
    else:
        st.success("No single skill would unlock more jobs at this match % — raise the Job Description Match % in the sidebar.")

# ----------------------------------------------------------------------------- 4 · fastest path

with st.container(border=True):
    section(4, f"Your fastest path — in {budget} days", "greedy search over the skill graph, on a scratch copy of you")
    if plan and plan["steps"]:
        total_after = plan["base"] + plan["total_unlocked"]
        prev = before.get("plan") if (before and before["sig"] != sig) else None
        d_days = (f'<span class="delta" style="color:#059669">{prev["total_days"]} → {plan["total_days"]} days</span>'
                  if prev and prev.get("steps") and prev["total_days"] != plan["total_days"] else "")
        st.markdown(f'<p class="big"><em>{plan["total_days"]} days</em> → +{plan["total_unlocked"]} jobs{d_days}</p>'
                    f'<p class="sub">{plan["base"]} → {total_after} matching jobs at ≥ {cutoff_pct}% JD match</p>', unsafe_allow_html=True)
        steps = "".join(
            f'<div class="step"><div class="day">Day {s["cumulative_days"]}</div><div class="sk">{s["skill"]}</div>'
            f'<div class="m">+{s["unlocks"]} jobs · {s["days"]} days</div><div class="c">{course_links(s["skill"], 1) or (s["course"] or "")}</div></div>'
            for s in plan["steps"])
        st.markdown(f'<div class="path">{steps}</div>', unsafe_allow_html=True)
    else:
        st.info("Nothing learnable in this many days gets you more matches — give it more days or lower the Job Description Match %.")

# ----------------------------------------------------------------------------- 5 · agent

with st.container(border=True):
    section(5, "Tell the agent what you learned", "ai& tool calling → one new edge in Neo4j → everything above re-plans")
    if delta_jobs:
        new = ", ".join(sorted(st.session_state.get("highlight", set())))
        st.markdown(f'<div class="banner">🧠 Memory updated: <b>{new}</b> → {"+" if delta_jobs > 0 else ""}{delta_jobs} jobs unlocked. '
                    f'Leaderboard, path and scores re-planned from the new edge.</div>', unsafe_allow_html=True)
    chat = st.session_state.setdefault("chat", [])
    for role, text in chat[-6:]:
        with st.chat_message(role):
            st.markdown(text)
    fb1, fb2 = st.columns([1, 3])
    if top and fb1.button(f"✅ I finished {short(top['skill'])}", help="Fallback: records the top skill without the LLM", width="stretch"):
        apply_skill(top["skill"], max(2, top["level"] or 2), source="button")
        chat.append(("user", f"I finished the {top['skill']} course."))
        chat.append(("assistant", f"Recorded **{top['skill']}** (level 2) in your memory graph. Re-planning…"))
        st.rerun()
    fb2.markdown('<p class="note" style="margin-top:10px">Try: “I just finished a GitHub Actions CI/CD course and set up pipelines for two repos”</p>',
                 unsafe_allow_html=True)

if msg := st.chat_input("Tell the agent what you learned…"):
    chat = st.session_state.setdefault("chat", [])
    chat.append(("user", msg))
    context = (f"Candidate skills: {', '.join(s['name'] + ' L' + str(s['level']) for s in skills)}. "
               f"Top learn-next options: {', '.join(r['skill'] + ' (+' + str(r['unlocks']) + ' jobs, ' + str(r['days']) + 'd)' for r in lb[:5])}.")
    try:
        with st.spinner(f"{model_label} agent thinking…"):
            calls, text, secs = llm.agent_turn(msg, context, model=model)
        st.session_state["agent_secs"] = secs
    except Exception as e:  # noqa: BLE001
        calls, text, secs = [], f"⚠️ ai& unavailable ({type(e).__name__}). Use the ✅ button to record the skill.", 0.0
    for c in calls:
        apply_skill(c["skill"], c["level"], source="agent")
    if calls:
        text = (text + "\n\n" if text else "") + "🧠 Recorded in Neo4j: " + ", ".join(
            f'**{c["skill"]}** (level {c["level"]})' for c in calls) + f" · {secs:.1f}s on ai&"
    chat.append(("assistant", text or "(no reply)"))
    st.rerun()

# ----------------------------------------------------------------------------- sidebar status (filled last)

ms = graph.LAST_MS
TIPS = {
    "fit": "Time for Neo4j to score every posting against your skill graph — one Cypher query over all REQUIRES edges.",
    "leaderboard": "Time for the learn-next query: which single skill would push the most postings over your cutoff, and at what cost in days.",
    "write": "Time to record a newly learned skill as a HAS edge in your memory graph — the agent's only write.",
    "extract": "Time for the ai& model to read your resume and structure it into the 44-skill taxonomy with levels.",
    "model": "The open-weight model serving extraction and the agent on ai& inference, hosted in Japan.",
}
row = lambda key, lbl, val: f'<div class="row" title="{TIPS[key]}"><span class="lbl">{lbl}</span><span class="ok">{val}</span></div>'
status.markdown(
    '<div class="status">'
    + row("fit", f"Neo4j · fit ×{len(fits)}", f'{ms.get("job_fits", 0):.0f} ms')
    + row("leaderboard", "Neo4j · leaderboard", f'{ms.get("leaderboard", 0):.0f} ms')
    + row("write", "Neo4j · memory write", f'{ms.get("mark_skill", 0):.0f} ms')
    + row("extract", "ai& · resume → skills", f'{(st.session_state.get("extract_secs") or 0):.1f} s')
    + row("model", "model", model_label)
    + '</div>', unsafe_allow_html=True)
